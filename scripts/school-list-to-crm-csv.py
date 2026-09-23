#!/usr/bin/env python3
"""
School prospect workbook -> one CSV for the CRM's Schools upload.

Run: python3 scripts/school-list-to-crm-csv.py <input.xlsx> <output.csv> [--report report.md]
     python3 scripts/school-list-to-crm-csv.py --self-test

Needs openpyxl (pip install openpyxl). Nothing here touches the app or the
database: it reads the workbook and writes two files.

The workbook is the athletic-programs prospect list: three list sheets
('Public High Schools', 'Private High Schools', 'Colleges') with a title block
above the header, a 'Summary' sheet of control totals and a 'Methodology' sheet.
The CRM stores a school as a row in the SAME fct_crm_companies blob as the
companies, so the CSV carries only what that row can hold, under headers the
upload modal normalises to the field names (see HEADER below).

What it does to the data, and why:

  * School names come from NCES/IPEDS title-cased, which leaves artifacts:
    'Mckinley', "Mary'S", 'Suny', 'Canon-Mcmillan Shs', 'University Of ...'.
    Those are fixed. SPELLING is not: 'Technolgy' is the official NCES spelling
    and a "fix" would stop the name matching the source it will be checked
    against. Every changed name is listed in the report for a human to read.
  * Addresses are re-cased only when the source is ALL CAPS (the private-school
    and some public rows). A mixed-case address was typed by someone; it stays.
  * ZIPs arrive as 9-digit strings with no hyphen ('215022596'); they go out as
    '21502-2596'. Every value is written as text so a leading zero survives.
  * Institution IDs are three different schemes (12-digit NCES, 'A2101928' PSS,
    6-digit IPEDS). They go out exactly as the workbook has them — never through
    int(), which is how a leading zero disappears.
  * The colleges' 'Athletics URL' is the same generic Equity in Athletics page
    on every row, so it is dropped; the high schools' MaxPreps link is kept.

The output is UTF-8 WITH a BOM (Excel otherwise reads it as ANSI), CRLF line
endings, RFC 4180 quoting, and NO newline inside any value. That last one is
not style: the tracker's _crmParseCSV splits the text into lines BEFORE it
looks at quotes, so a quoted newline turns one school into a broken row plus a
junk row. The writer refuses to emit one.

Counts are tied out against the workbook's own Summary sheet. A mismatch still
writes both files, so the report can say what is off, but exits 1: an import
that does not tie to its source should not be uploaded on a green run. Know what
that proves: the Summary cells are COUNTIF formulas over these same sheets, so
a tie means the conversion dropped, added and relabelled nothing — not that the
research behind the list is complete. The script reads their CACHED values; a
workbook last saved by a tool that does not recalculate has none, and every
check then fails rather than passing on blanks.
"""
import argparse
import csv
import datetime
import io
import os
import re
import sys
import unicodedata
from collections import Counter, OrderedDict, defaultdict
from decimal import Decimal, ROUND_HALF_UP

# (sheet name, CRM Type). Order here is the order rows are written.
SHEETS = [
    ('Public High Schools',  'High School'),
    ('Private High Schools', 'High School'),
    ('Colleges',             'College / University'),
]

# The workbook's list headers. Columns are found by these names, not by
# position, so a reordered or widened sheet still converts correctly — and a
# renamed one fails loudly instead of shifting every value one column over.
SRC_COLS = ['School', 'City', 'County', 'State', 'Public/Private', 'Boundary status',
            'Miles to boundary', 'Address', 'ZIP', 'Enrollment', 'Athletics classification',
            'Athletics evidence', 'Athletics URL', 'Institution ID', 'Source year',
            'Athletics match review']

# Output columns, in upload order. The upload modal normalises each header with
# h.toLowerCase().replace(/[^a-z0-9]+/g,'_') — 'Public / Private' becomes
# public_private, 'Mi to Line' mi_to_line, 'NCES ID' nces_id.
HEADER = ['Tag', 'Lead Contact', 'School', 'Type', 'Public / Private', 'Athletics',
          'Enrollment', 'Territory', 'Mi to Line', 'City', 'County', 'State', 'Zip',
          'Address', 'Phone', 'Email Domain', 'Website', 'Athletics Page',
          'Personal Interest', 'Turf Product', 'Notes', 'NCES ID']

TERRITORIES = ('Inside', 'On/near boundary')
REVIEW_NAMING = 'Review school/profile naming'
NOTE_REVIEW_NAMING = 'Check the MaxPreps profile - its name does not match the school'


# ── name cleaning ────────────────────────────────────────────────────────────

# Abbreviations that title-casing turned into words ('Jshs', 'Suny').
# UMS is Farrell's 'HS/UMS' (upper middle school). MS appears only as 'Ms/Hs'.
NAME_UPPER = {'HS', 'SHS', 'JSHS', 'JHS', 'MS', 'UMS', 'CS', 'SUNY', 'CUNY', 'STEM',
              'II', 'III', 'IV'}
# Lower-case unless first word. 'de' is the particle in 'Martin de Porres' and
# 'Eugenio Maria de Hostos'; a leading 'De' ('DeSales') is the first word and
# is left alone by the same rule.
SMALL_WORDS = {'of', 'the', 'and', 'at', 'for', 'in', 'on', 'de'}
# Names whose own spelling has an inner capital that title-casing flattened.
# Kept to names confirmed in this list; each change lands in the report.
CAMEL_CASE = {'dubois': 'DuBois', 'leboeuf': 'LeBoeuf', 'lasalle': 'LaSalle',
              'labrae': 'LaBrae'}

_WS = re.compile(r'\s+')


def squash(s):
    """Trim, and fold every whitespace run — including any newline — to one space."""
    return _WS.sub(' ', '' if s is None else str(s)).strip()


def _mc(word):
    # 'Mckinley' -> 'McKinley'. Needs a letter after 'Mc' ('Mc' alone is left).
    return re.sub(r'^Mc([a-z])', lambda m: 'Mc' + m.group(1).upper(), word)


def clean_name(raw):
    """Undo title-case artifacts in a school name. Never changes spelling."""
    s = squash(raw)
    # Split into words, keeping the separators, so 'Canon-Mcmillan' and
    # 'Ms/Hs' are handled a part at a time and put back exactly as they were.
    parts = re.split(r'(\s+|-|/)', s)
    out, first, prev_sep = [], True, ''
    for p in parts:
        if p == '':
            continue
        if re.fullmatch(r'\s+|-|/', p):
            out.append(p)
            prev_sep = p
            continue
        # Punctuation hugging the word ('Jr.', '(Hs)') is kept outside the match.
        m = re.match(r"^([^A-Za-z0-9']*)(.*?)([^A-Za-z0-9']*)$", p)
        lead, core, trail = m.group(1), m.group(2), m.group(3)
        low = core.lower()
        if core.upper() in NAME_UPPER:
            core = core.upper()
        elif low in CAMEL_CASE:
            core = CAMEL_CASE[low]
        elif not first and low in SMALL_WORDS and prev_sep not in ('-', '/'):
            core = low
        else:
            core = _mc(core)
            # "John'S" -> "John's". Only a trailing 'S: "D'Youville" is a name.
            core = re.sub(r"'S$", "'s", core)
        out.append(lead + core + trail)
        first, prev_sep = False, ''
    return ''.join(out)


# ── address cleaning ─────────────────────────────────────────────────────────

# Tokens that stay upper case in an address: directionals, box/route prefixes.
ADDR_UPPER = {'N', 'S', 'E', 'W', 'NE', 'NW', 'SE', 'SW', 'PO', 'RR', 'US', 'SR',
              'CR', 'TR', 'HC', 'II', 'III', 'IV'}
ADDR_SMALL = {'of', 'the', 'and', 'at', 'for', 'in', 'on'}


def _cap(word):
    # One address word out of ALL CAPS: 'MCMULLEN' -> 'McMullen',
    # "O'NEIL" -> "O'Neil", 'WILKES-BARRE' -> 'Wilkes-Barre'.
    pieces = []
    for piece in word.split('-'):
        w = piece[:1].upper() + piece[1:].lower()
        w = re.sub(r"^([A-Z])'([a-z])", lambda m: m.group(1) + "'" + m.group(2).upper(), w)
        pieces.append(_mc(w))
    return '-'.join(pieces)


def clean_addr(raw):
    """Title-case an ALL-CAPS street address; leave a mixed-case one as typed."""
    s = squash(raw)
    if not s or s != s.upper() or not re.search(r'[A-Z]', s):
        return s
    out = []
    for i, tok in enumerate(s.split(' ')):
        m = re.match(r'^([^A-Z0-9]*)(.*?)([^A-Z0-9]*)$', tok)
        lead, core, trail = m.group(1), m.group(2), m.group(3)
        if core in ADDR_UPPER or len(core) <= 1:
            pass                                        # 'SW', 'PO', unit 'B'
        elif re.fullmatch(r'\d+(ST|ND|RD|TH)', core):
            core = core.lower()                         # '1ST' -> '1st'
        elif re.search(r'\d', core):
            pass                                        # '12A', '31-35', '#29'
        elif i > 0 and core.lower() in ADDR_SMALL:
            core = core.lower()                         # 'FOOT OF TEN RD'
        else:
            core = _cap(core)
        out.append(lead + core + trail)
    return ' '.join(out)


# ── the other fields ─────────────────────────────────────────────────────────

def clean_zip(raw):
    """'215022596' -> '21502-2596'. Returns (zip, problem-or-None)."""
    if raw is None or str(raw).strip() == '':
        return '', None
    if isinstance(raw, (int, float)):
        # A numeric cell has already lost its leading zeros; put them back.
        n = int(raw)
        return clean_zip(str(n).zfill(5 if n < 100000 else 9))
    digits = re.sub(r'\D', '', str(raw))
    if len(digits) == 5:
        return digits, None
    if len(digits) == 9:
        return digits[:5] + '-' + digits[5:], None
    # A text ZIP of any other length is damaged, and padding it would be a guess.
    return '', 'ZIP %r is not 5 or 9 digits; left blank' % (raw,)


def clean_county(raw):
    """'Allegany' -> 'Allegany County'. The private sheet omits the word."""
    c = squash(raw)
    if not c or re.search(r'\b(county|city)$', c, re.I):
        return c
    return c + ' County'


def clean_miles(raw):
    """Always one decimal: 2 -> '2.0'. Decimal, so 0.05 rounds up, not to 0.0."""
    if raw is None or str(raw).strip() == '':
        return ''
    d = Decimal(str(raw).strip()).quantize(Decimal('0.1'), rounding=ROUND_HALF_UP)
    return '%.1f' % d


def clean_int(raw):
    """Enrollment as a whole number, blank when missing."""
    if raw is None or str(raw).strip() == '':
        return ''
    if isinstance(raw, float):
        return str(int(round(raw)))
    return str(int(Decimal(re.sub(r'[,\s]', '', str(raw)))))


def ascii_only(s):
    # Notes quote other schools' names; keep them plain for any downstream tool.
    return unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode('ascii')


def url_key(u):
    u = squash(u).lower()
    return u if u.endswith('/') else u + '/'


# ── reading the workbook ─────────────────────────────────────────────────────

def read_list_sheet(ws):
    """Rows of one list sheet as dicts keyed by SRC_COLS, plus the sheet row number.

    The header is found by content (the row whose first cell is 'School') rather
    than assumed to be row 5, because the title block above it is prose that
    someone will eventually add a line to.
    """
    header_row, cols = None, None
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True), 1):
        if row and squash(row[0]) == 'School':
            header_row = r_idx
            cols = {squash(h): i for i, h in enumerate(row) if h is not None}
            break
    if header_row is None:
        raise SystemExit("%s: no header row starting with 'School' in the first 30 rows" % ws.title)
    missing = [c for c in SRC_COLS if c not in cols]
    if missing:
        raise SystemExit('%s: header is missing %s' % (ws.title, ', '.join(missing)))
    rows = []
    for r_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True),
                                header_row + 1):
        if not any(v is not None and str(v).strip() != '' for v in row):
            continue
        rec = {c: (row[cols[c]] if cols[c] < len(row) else None) for c in SRC_COLS}
        rec['_row'] = r_idx
        rows.append(rec)
    return rows


def read_summary(wb):
    """The Summary sheet's control totals, found by their labels."""
    if 'Summary' not in wb.sheetnames:
        return None
    ws = wb['Summary']
    out = {'by_state': OrderedDict()}
    in_state_table, state_cols = False, None
    for row in ws.iter_rows(values_only=True):
        label = squash(row[0]) if row and row[0] is not None else ''
        if label == 'Total programs':
            out['total'] = row[1]
        elif label == 'Clearly inside':
            out['Inside'] = row[1]
        elif label == 'On/near boundary':
            out['On/near boundary'] = row[1]
        elif label == 'State':
            in_state_table = True
            state_cols = [squash(h) for h in row]
        elif in_state_table and label:
            vals = dict(zip(state_cols, row))
            entry = {
                'Public High Schools': vals.get('Public high schools'),
                'Private High Schools': vals.get('Private high schools'),
                'Colleges': vals.get('Colleges'),
                'Total': vals.get('Total'),
            }
            if label == 'Total':
                out['by_sheet'] = entry
                in_state_table = False
            else:
                out['by_state'][label] = entry
    return out


# ── conversion ───────────────────────────────────────────────────────────────

def convert(wb):
    """Returns (records, problems). A record is the output row plus its source."""
    records, problems = [], []
    for sheet, type_label in SHEETS:
        if sheet not in wb.sheetnames:
            raise SystemExit('workbook has no %r sheet' % sheet)
        is_college = type_label.startswith('College')
        for src in read_list_sheet(wb[sheet]):
            where = '%s row %d' % (sheet, src['_row'])
            nces = src['Institution ID']
            if not isinstance(nces, str):
                # A numeric ID cell may already have dropped a leading zero.
                problems.append('%s: Institution ID %r is not text; check it for lost leading zeros'
                                % (where, nces))
            territory = squash(src['Boundary status'])
            if territory not in TERRITORIES:
                problems.append('%s: Boundary status %r is not Inside or On/near boundary'
                                % (where, territory))
            sector = squash(src['Public/Private'])
            if sector not in ('Public', 'Private'):
                problems.append('%s: Public/Private %r' % (where, sector))
            zip_, zip_problem = clean_zip(src['ZIP'])
            if zip_problem:
                problems.append('%s: %s' % (where, zip_problem))
            miles = clean_miles(src['Miles to boundary'])
            if not miles:
                problems.append('%s: no Miles to boundary' % where)
            url = squash(src['Athletics URL'])
            if not is_college and url and 'maxpreps.com/' not in url.lower():
                problems.append('%s: athletics URL is not a MaxPreps profile: %s' % (where, url))
            out = OrderedDict((h, '') for h in HEADER)
            out.update({
                'School': clean_name(src['School']),
                'Type': type_label,
                'Public / Private': sector,
                'Athletics': squash(src['Athletics classification']) if is_college else '',
                'Enrollment': clean_int(src['Enrollment']),
                'Territory': territory,
                'Mi to Line': miles,
                'City': squash(src['City']),
                'County': clean_county(src['County']),
                'State': squash(src['State']).upper(),
                'Zip': zip_,
                'Address': clean_addr(src['Address']),
                'Athletics Page': '' if is_college else url,
                'NCES ID': squash(nces),
            })
            records.append({'sheet': sheet, 'src': src, 'out': out, 'flags': []})

    # Flag (a): the source's own "the profile's name doesn't match" marker.
    for rec in records:
        if squash(rec['src']['Athletics match review']) == REVIEW_NAMING:
            rec['flags'].append(('review', NOTE_REVIEW_NAMING))
        elif squash(rec['src']['Athletics match review']):
            problems.append('%s row %d: unrecognised Athletics match review %r'
                            % (rec['sheet'], rec['src']['_row'], rec['src']['Athletics match review']))

    # Flag (b): one MaxPreps profile on several rows — usually an elementary or
    # middle campus listed beside its high school, or two campuses of one
    # program. Across sheets it is more often a wrong match (a public school
    # given a private school's profile), which the report calls out.
    by_url = defaultdict(list)
    for rec in records:
        if rec['out']['Athletics Page']:
            by_url[url_key(rec['out']['Athletics Page'])].append(rec)
    for group in by_url.values():
        if len(group) < 2:
            continue
        for rec in group:
            others = [o['out']['School'] for o in group if o is not rec]
            names = others[0] if len(others) == 1 else ', '.join(others[:-1]) + ' and ' + others[-1]
            rec['flags'].append(('shared', 'Shares a MaxPreps profile with ' + names))

    for rec in records:
        rec['out']['Notes'] = ascii_only('; '.join(text for _, text in rec['flags']))
    return records, problems


def write_csv(records, path):
    buf = io.StringIO()
    # QUOTE_MINIMAL quotes exactly the fields RFC 4180 requires (comma, quote,
    # line break) and doubles inner quotes.
    w = csv.writer(buf, lineterminator='\r\n', quoting=csv.QUOTE_MINIMAL)
    w.writerow(HEADER)
    for rec in records:
        row = [rec['out'][h] for h in HEADER]
        for h, v in zip(HEADER, row):
            if '\n' in v or '\r' in v:
                # squash() should make this impossible; the check is the promise.
                raise SystemExit('refusing to write a line break inside %s for %s'
                                 % (h, rec['out']['School']))
        w.writerow(row)
    with open(path, 'w', encoding='utf-8-sig', newline='') as f:
        f.write(buf.getvalue())


# ── the report ───────────────────────────────────────────────────────────────

def _md(v):
    return str(v).replace('|', '\\|') if v not in (None, '') else ''


def _table(headers, rows):
    lines = ['| ' + ' | '.join(headers) + ' |', '|' + '---|' * len(headers)]
    lines += ['| ' + ' | '.join(_md(c) for c in r) + ' |' for r in rows]
    return '\n'.join(lines)


def _name_change_kind(before, after):
    """Why a name changed, in words a reviewer can scan."""
    b, a = squash(before), after
    kinds = []
    if squash(before) != str(before or ''):
        kinds.append('spacing')
    bw, aw = re.split(r'(\s+|-|/)', b), re.split(r'(\s+|-|/)', a)
    if len(bw) != len(aw):
        return ', '.join(kinds + ['other'])
    for x, y in zip(bw, aw):
        if x == y:
            continue
        core = re.sub(r"[^A-Za-z']", '', y)
        if core.upper() in NAME_UPPER and core == core.upper():
            kinds.append('abbreviation')
        elif core.lower() in CAMEL_CASE:
            kinds.append('inner capital')
        elif core in SMALL_WORDS:
            kinds.append('small word')
        elif core.startswith('Mc'):
            kinds.append('Mc')
        elif core.endswith("'s"):
            kinds.append('possessive')
        else:
            kinds.append('other')
    return ', '.join(OrderedDict.fromkeys(kinds))


def build_report(records, problems, summary, in_path, out_path):
    L = []
    add = L.append
    n = len(records)
    add('# Schools CRM import report')
    add('')
    add('- Source: `%s`' % os.path.basename(in_path))
    add('- Output: `%s` (%d data rows, %d columns)' % (os.path.basename(out_path), n, len(HEADER)))
    add('- Generated: %s by `scripts/school-list-to-crm-csv.py`'
        % datetime.datetime.now().strftime('%Y-%m-%d %H:%M'))
    add('')

    # Tie-out against the Summary sheet.
    ties = []   # (check, expected, actual)
    by_sheet = Counter(r['sheet'] for r in records)
    by_state_sheet = Counter((r['out']['State'], r['sheet']) for r in records)
    by_state = Counter(r['out']['State'] for r in records)
    by_terr = Counter(r['out']['Territory'] for r in records)
    add('## Row counts tied to the Summary sheet')
    add('')
    add('The Summary cells are COUNTIF formulas over the three list sheets, so a match proves '
        'the conversion dropped, added and relabelled nothing. It does not prove the research '
        'behind the list is complete.')
    add('')
    if summary is None:
        add('**The workbook has no Summary sheet; nothing to tie out against.**')
    else:
        ties.append(('Total rows', summary.get('total'), n))
        for t in TERRITORIES:
            ties.append(('Territory: ' + t, summary.get(t), by_terr.get(t, 0)))
        for sheet, _ in SHEETS:
            ties.append(('Sheet: ' + sheet, (summary.get('by_sheet') or {}).get(sheet), by_sheet[sheet]))
        states = sorted(set(summary['by_state']) | set(by_state))
        for st in states:
            exp = summary['by_state'].get(st, {})
            ties.append(('State %s total' % st, exp.get('Total'), by_state.get(st, 0)))
            for sheet, _ in SHEETS:
                ties.append(('State %s, %s' % (st, sheet), exp.get(sheet), by_state_sheet[(st, sheet)]))
        add(_table(['Check', 'Summary sheet', 'CSV', 'Result'],
                   [(c, e, a, 'match' if e == a else '**MISMATCH**') for c, e, a in ties]))
        bad = [t for t in ties if t[1] != t[2]]
        add('')
        add('**All %d checks match.**' % len(ties) if not bad
            else '**%d of %d checks do NOT match. Do not upload until they are explained.**'
                 % (len(bad), len(ties)))
    add('')
    add('Territory by sheet:')
    add('')
    add(_table(['Sheet', 'Inside', 'On/near boundary', 'Total'],
               [(s, sum(1 for r in records if r['sheet'] == s and r['out']['Territory'] == 'Inside'),
                 sum(1 for r in records if r['sheet'] == s and r['out']['Territory'] == 'On/near boundary'),
                 by_sheet[s]) for s, _ in SHEETS]))
    add('')
    # The Summary defines Inside as "more than 25 miles from the line" and
    # On/near as "within 25 miles". Check the labels agree with the miles.
    odd_terr = [r for r in records if r['out']['Mi to Line'] and (
        (r['out']['Territory'] == 'Inside' and Decimal(r['out']['Mi to Line']) <= 25) or
        (r['out']['Territory'] == 'On/near boundary' and Decimal(r['out']['Mi to Line']) > 25))]
    add('Territory vs miles: %s' % ('every Inside row is more than 25 mi from the line and every '
                                     'On/near row is within 25 mi.' if not odd_terr else
                                     '**%d rows disagree with the 25-mile rule** (listed under anomalies).'
                                     % len(odd_terr)))
    add('')

    # Flags.
    flag_counts = Counter(kind for r in records for kind, _ in r['flags'])
    flagged = [r for r in records if r['flags']]
    add('## Notes flags')
    add('')
    add(_table(['Flag', 'Rows'], [
        ('Check the MaxPreps profile - its name does not match the school', flag_counts['review']),
        ('Shares a MaxPreps profile with another row', flag_counts['shared']),
        ('Rows with any flag', len(flagged)),
    ]))
    add('')
    add(_table(['School', 'City', 'State', 'Sheet', 'Notes'],
               [(r['out']['School'], r['out']['City'], r['out']['State'], r['sheet'], r['out']['Notes'])
                for r in flagged]))
    add('')

    # Name changes.
    changes = [(r, squash(r['src']['School']) != r['out']['School'] or r['src']['School'] != r['out']['School'])
               for r in records]
    changes = [r for r, c in changes if c]
    abbrev_only = [r for r in changes
                   if _name_change_kind(r['src']['School'], r['out']['School']) == 'abbreviation']
    other = [r for r in changes if r not in abbrev_only]
    add('## School name changes (%d of %d names)' % (len(changes), n))
    add('')
    add('Spelling is never changed. Every change is below; the %d that only upper-case an '
        'abbreviation (Hs -> HS, Jshs -> JSHS, Suny -> SUNY, ...) are listed last.' % len(abbrev_only))
    add('')
    add('### Changes other than abbreviations (%d) - read these' % len(other))
    add('')
    add(_table(['Before', 'After', 'Why', 'City', 'State'],
               [(repr(r['src']['School']) if squash(r['src']['School']) != r['src']['School']
                 else r['src']['School'], r['out']['School'],
                 _name_change_kind(r['src']['School'], r['out']['School']),
                 r['out']['City'], r['out']['State']) for r in other]))
    add('')
    add('### Abbreviation-only changes (%d)' % len(abbrev_only))
    add('')
    add(_table(['Before', 'After', 'City', 'State'],
               [(r['src']['School'], r['out']['School'], r['out']['City'], r['out']['State'])
                for r in abbrev_only]))
    add('')

    # Addresses.
    addr_changed = [r for r in records if squash(r['src']['Address']) and
                    squash(r['src']['Address']) != r['out']['Address']]
    ws_only = [r for r in records if r['src']['Address'] and
               str(r['src']['Address']) != squash(r['src']['Address'])]
    add('## Addresses')
    add('')
    add('%d addresses were ALL CAPS and were re-cased; %d had stray whitespace collapsed. '
        'Mixed-case addresses are otherwise exactly as the source.' % (len(addr_changed), len(ws_only)))
    add('')
    add('<details><summary>All re-cased addresses</summary>')
    add('')
    add(_table(['Before', 'After', 'School', 'State'],
               [(r['src']['Address'], r['out']['Address'], r['out']['School'], r['out']['State'])
                for r in addr_changed]))
    add('')
    add('</details>')
    add('')
    missing = [r for r in records if not r['out']['Address'] or not r['out']['Zip']]
    add('### Rows missing an address or ZIP (%d)' % len(missing))
    add('')
    miss_by_sheet = Counter(r['sheet'] for r in missing)
    add(', '.join('%s: %d of %d' % (s, miss_by_sheet[s], by_sheet[s]) for s, _ in SHEETS) + '.')
    add('')
    add(_table(['School', 'City', 'State', 'Sheet', 'Missing'],
               [(r['out']['School'], r['out']['City'], r['out']['State'], r['sheet'],
                 ' and '.join(x for x, v in (('address', r['out']['Address']), ('ZIP', r['out']['Zip']))
                              if not v)) for r in missing]))
    add('')

    # Enrollment.
    add('## Enrollment coverage')
    add('')
    add(_table(['Sheet', 'With enrollment', 'Rows'],
               [(s, sum(1 for r in records if r['sheet'] == s and r['out']['Enrollment']), by_sheet[s])
                for s, _ in SHEETS] +
               [('All', sum(1 for r in records if r['out']['Enrollment']), n)]))
    add('')

    # Anomalies.
    add('## Anomalies for a human to look at')
    add('')
    anomalies = []
    for p in problems:
        anomalies.append(('Conversion problem', p))
    for r in odd_terr:
        anomalies.append(('Territory disagrees with miles', '%s (%s, %s): %s at %s mi'
                          % (r['out']['School'], r['out']['City'], r['out']['State'],
                             r['out']['Territory'], r['out']['Mi to Line'])))
    # Same name twice in one state.
    by_name_state = defaultdict(list)
    for r in records:
        by_name_state[(r['out']['School'].lower(), r['out']['State'])].append(r)
    for (_, st), group in sorted(by_name_state.items(), key=lambda kv: (kv[0][1], kv[0][0])):
        if len(group) > 1:
            cities = [g['out']['City'] for g in group]
            same_city = len(set(cities)) < len(cities)
            anomalies.append(('Same name in one state' + (' AND city' if same_city else ''),
                              '%s (%s): %s' % (group[0]['out']['School'], st,
                                               ', '.join('%s [%s]' % (g['out']['City'], g['out']['NCES ID'])
                                                         for g in group))))
    # Same street address on several rows.
    by_addr = defaultdict(list)
    for r in records:
        if r['out']['Address']:
            by_addr[(r['out']['Address'].lower(), r['out']['Zip'][:5])].append(r)
    for (addr, _), group in by_addr.items():
        if len(group) > 1:
            anomalies.append(('Same address on several rows', '%s, %s: %s'
                              % (group[0]['out']['Address'], group[0]['out']['City'],
                                 '; '.join(g['out']['School'] for g in group))))
    # A MaxPreps profile shared across the public and private sheets is a
    # public school carrying a private school's profile (or the reverse).
    by_url = defaultdict(list)
    for r in records:
        if r['out']['Athletics Page']:
            by_url[url_key(r['out']['Athletics Page'])].append(r)
    for group in by_url.values():
        if len(group) > 1 and len({g['sheet'] for g in group}) > 1:
            anomalies.append(('Profile shared across public/private', '%s -> %s; probably a wrong match'
                              % (' / '.join('%s (%s)' % (g['out']['School'], g['sheet']) for g in group),
                                 group[0]['out']['Athletics Page'])))
    for r in records:
        o = r['out']
        url = o['Athletics Page']
        m = re.match(r'https?://(?:www\.)?maxpreps\.com/([a-z]{2})/', url, re.I)
        if url and m and m.group(1).upper() != o['State']:
            anomalies.append(('MaxPreps state differs', '%s (%s): %s' % (o['School'], o['State'], url)))
        if r['sheet'] != 'Colleges' and re.search(r'\b(elementary|preschool)\b|\bmiddle school$|^\w+ middle school$',
                                                  o['School'], re.I) \
                and not re.search(r'high|senior|secondary|\bhs\b|jr|sr', o['School'], re.I):
            anomalies.append(('Name does not read as a high school',
                              '%s (%s, %s) - NCES lists it with grade 12; confirm it fields varsity teams'
                              % (o['School'], o['City'], o['State'])))
        # 'One College Ave', 'PO Box 6201' and grid-style 'S 4432 Bay View Rd'
        # are real mailing addresses. What is left has no house number at all:
        # a building name, an office, or the school's own name typed as the address.
        if o['Address'] and not re.search(r'\d', o['Address']) \
                and not re.match(r'(one|two|three|four|five)\b', o['Address'], re.I):
            anomalies.append(('Address has no house number', '%s (%s, %s): %r'
                              % (o['School'], o['City'], o['State'], o['Address'])))
        elif o['Address'] and re.search(r',| - |\b(bldg|office|floor|suite)\b'
                                        r'|\b(rd|dr|ave|road|drive|avenue|blvd)\b.*\d', o['Address'], re.I):
            anomalies.append(('Address carries office/suite text', '%s (%s, %s): %r'
                              % (o['School'], o['City'], o['State'], o['Address'])))
        if re.search(r'\S- | -\S', o['School']):
            anomalies.append(('Odd spacing around a hyphen (kept as source)', o['School']))
        if o['Mi to Line'] == '0.0':
            anomalies.append(('On the line', '%s (%s, %s) is 0.0 mi from the boundary'
                              % (o['School'], o['City'], o['State'])))
        if o['Athletics'] == 'Other':
            anomalies.append(("College athletics classification is 'Other'", '%s (%s, %s)'
                              % (o['School'], o['City'], o['State'])))
        for h, v in o.items():
            if any(ord(ch) > 127 for ch in v):
                anomalies.append(('Non-ASCII text', '%s: %s = %r' % (o['School'], h, v)))
    # City goes out exactly as the source has it, so its title-case artifacts
    # stay too. Name them, so the CRM's city filter surprises nobody.
    city_forms = defaultdict(Counter)
    for r in records:
        city_forms[(re.sub(r'[\s.]', '', r['out']['City'].lower()), r['out']['State'])][r['out']['City']] += 1
    for (_, st), forms in sorted(city_forms.items(), key=lambda kv: (kv[0][1], kv[0][0])):
        names = sorted(forms)
        if len(forms) > 1:
            anomalies.append(('City spelled more than one way (kept as source)', '%s: %s' % (
                st, ' / '.join('%s (%d rows)' % (c, forms[c]) for c in names))))
        for c in names:
            if re.search(r'\bMc[a-z]', c):
                anomalies.append(('City has a title-case artifact (kept as source)',
                                  '%s, %s (%d rows)' % (c, st, forms[c])))
    # One kind at a time reads better than source order.
    kinds = list(OrderedDict.fromkeys(k for k, _ in anomalies))
    anomalies.sort(key=lambda a: kinds.index(a[0]))
    add(_table(['Kind', 'Detail'], anomalies) if anomalies else 'None found.')
    add('')
    return '\n'.join(L), ties


# ── self-test ────────────────────────────────────────────────────────────────

def self_test():
    cases = [
        (clean_name, 'Mckinley High School', 'McKinley High School'),
        (clean_name, 'Mcdaniel College', 'McDaniel College'),
        (clean_name, "Mount St. Mary'S University", "Mount St. Mary's University"),
        (clean_name, "St John'S Catholic Prep", "St John's Catholic Prep"),
        (clean_name, "D'Youville  University", "D'Youville University"),
        (clean_name, 'Suny College Of Technology At Alfred', 'SUNY College of Technology at Alfred'),
        (clean_name, 'Canon-Mcmillan Shs', 'Canon-McMillan SHS'),
        (clean_name, 'University Of Pittsburgh-Johnstown', 'University of Pittsburgh-Johnstown'),
        (clean_name, 'School 58-World Of Inquiry School', 'School 58-World of Inquiry School'),
        (clean_name, 'Clairton Ms/Hs', 'Clairton MS/HS'),
        (clean_name, 'Farrell Area Hs/Ums', 'Farrell Area HS/UMS'),
        (clean_name, 'Somerset Area Jr-Sr Hs', 'Somerset Area Jr-Sr HS'),
        (clean_name, 'Lakeview Middle-Hs', 'Lakeview Middle-HS'),
        (clean_name, 'Propel Cs-Braddock Hills', 'Propel CS-Braddock Hills'),
        (clean_name, 'The College Of Wooster', 'The College of Wooster'),
        (clean_name, 'Of Mice School', 'Of Mice School'),           # first word stays
        (clean_name, 'Lewis J Bennett High School Of Innovative Technolgy',
         'Lewis J Bennett High School of Innovative Technolgy'),   # spelling kept
        (clean_name, 'St Martin De Porres High School', 'St Martin de Porres High School'),
        (clean_name, 'Seton Lasalle Catholic High School', 'Seton LaSalle Catholic High School'),
        (clean_name, 'Pennsylvania State University-Penn State Dubois',
         'Pennsylvania State University-Penn State DuBois'),
        (clean_name, 'Palmyra-Macedon Senior High School', 'Palmyra-Macedon Senior High School'),
        (clean_name, 'Leetonia Jr./Sr. High School', 'Leetonia Jr./Sr. High School'),
        (clean_name, 'Mc Hs', 'Mc HS'),
        (clean_name, 'Pennsylvania State University-Penn State Fayette- Eberly',
         'Pennsylvania State University-Penn State Fayette- Eberly'),
        (clean_name, 'Beaver Co Christian School -Upper', 'Beaver Co Christian School -Upper'),
        (clean_name, 'The Bridgeport School District - High School',
         'The Bridgeport School District - High School'),
        (clean_name, 'A\nB', 'A B'),
        (clean_name, None, ''),
        (clean_addr, '14517 MCMULLEN HWY SW', '14517 McMullen Hwy SW'),
        (clean_addr, '200 1ST LN', '200 1st Ln'),
        (clean_addr, '4599 BURBANK RD B', '4599 Burbank Rd B'),
        (clean_addr, '1140 FOOT OF TEN RD', '1140 Foot of Ten Rd'),
        (clean_addr, 'PO BOX 350', 'PO Box 350'),
        (clean_addr, 'RR 1 BOX 22', 'RR 1 Box 22'),
        (clean_addr, '1 N 2ND ST STE A', '1 N 2nd St Ste A'),
        (clean_addr, '17 OLIVER ST #29', '17 Oliver St #29'),
        (clean_addr, '31-35 ELM ST', '31-35 Elm St'),
        (clean_addr, '12 WILKES-BARRE BLVD', '12 Wilkes-Barre Blvd'),
        (clean_addr, "9 O'NEIL AVE", "9 O'Neil Ave"),
        (clean_addr, 'S 4432 BAY VIEW RD', 'S 4432 Bay View Rd'),
        (clean_addr, '900 Seton Dr', '900 Seton Dr'),
        (clean_addr, '100 Dr. Nancy S. Grasmick', '100 Dr. Nancy S. Grasmick'),
        (clean_addr, '319 South Market  Street', '319 South Market Street'),
        (clean_addr, '12401 Willowbrook Rd SE', '12401 Willowbrook Rd SE'),
        (clean_addr, None, ''),
        (clean_addr, '1234', '1234'),
        (clean_zip, '215022596', ('21502-2596', None)),
        (clean_zip, '21502', ('21502', None)),
        (clean_zip, '21502-2596', ('21502-2596', None)),
        (clean_zip, '01234', ('01234', None)),
        (clean_zip, 1234, ('01234', None)),
        (clean_zip, 12345678, ('01234-5678', None)),
        (clean_zip, None, ('', None)),
        (clean_zip, '', ('', None)),
        (clean_county, 'Allegany', 'Allegany County'),
        (clean_county, 'Erie County', 'Erie County'),
        (clean_county, ' Erie ', 'Erie County'),
        (clean_county, 'Baltimore city', 'Baltimore city'),
        (clean_county, None, ''),
        (clean_miles, 45.9, '45.9'),
        (clean_miles, 2, '2.0'),
        (clean_miles, 0, '0.0'),
        (clean_miles, 0.05, '0.1'),
        (clean_miles, 12.25, '12.3'),
        (clean_miles, None, ''),
        (clean_int, 219, '219'),
        (clean_int, 219.0, '219'),
        (clean_int, '1,604', '1604'),
        (clean_int, None, ''),
    ]
    bad = 0
    for fn, arg, want in cases:
        got = fn(arg)
        if got != want:
            bad += 1
            print('  FAIL %s(%r) = %r, want %r' % (fn.__name__, arg, got, want))
    # A bad text ZIP is reported, not padded into something plausible.
    z, why = clean_zip('2150')
    if z != '' or not why:
        bad += 1
        print('  FAIL clean_zip(%r) should be blank with a problem, got %r %r' % ('2150', z, why))
    print('self-test: %d cases, %d failed' % (len(cases) + 1, bad))
    return 1 if bad else 0


# ── main ─────────────────────────────────────────────────────────────────────

def main(argv):
    ap = argparse.ArgumentParser(description='Convert the school prospect workbook to a CRM Schools CSV.')
    ap.add_argument('input', nargs='?', help='the .xlsx workbook')
    ap.add_argument('output', nargs='?', help='the .csv to write')
    ap.add_argument('--report', help='also write a markdown review report here')
    ap.add_argument('--self-test', action='store_true', help='run the cleaning-rule tests and exit')
    args = ap.parse_args(argv)
    if args.self_test:
        return self_test()
    if not args.input or not args.output:
        ap.error('input and output are required')
    try:
        import openpyxl
    except ImportError:
        raise SystemExit('openpyxl is required: pip install openpyxl')

    # data_only: the cached values, not formulas. read_only is NOT used — it
    # can report a stale sheet dimension and stop short of the last rows.
    wb = openpyxl.load_workbook(args.input, data_only=True)
    records, problems = convert(wb)
    summary = read_summary(wb)

    ids = Counter(r['out']['NCES ID'] for r in records)
    dupes = [i for i, c in ids.items() if c > 1]
    if dupes:
        problems.append('Institution ID repeated: ' + ', '.join(dupes))
    blank_ids = sum(1 for r in records if not r['out']['NCES ID'])
    if blank_ids:
        problems.append('%d rows have no Institution ID' % blank_ids)

    write_csv(records, args.output)
    report, ties = build_report(records, problems, summary, args.input, args.output)
    if args.report:
        with open(args.report, 'w', encoding='utf-8', newline='\n') as f:
            f.write(report)

    mismatched = [t for t in ties if t[1] != t[2]]
    print('wrote %d rows to %s' % (len(records), args.output))
    if args.report:
        print('wrote report to %s' % args.report)
    print('flags: %d review-naming, %d shared-profile'
          % (sum(1 for r in records for k, _ in r['flags'] if k == 'review'),
             sum(1 for r in records for k, _ in r['flags'] if k == 'shared')))
    for p in problems:
        print('  !! ' + p)
    if summary is None:
        print('  !! no Summary sheet: counts NOT tied out')
        return 1
    if mismatched:
        for c, e, a in mismatched:
            print('  !! %s: Summary says %s, CSV has %s' % (c, e, a))
        if any(e is None for _, e, _ in mismatched):
            print('  (a blank Summary value is usually a formula with no cached result:'
                  ' open the workbook in Excel, let it calculate, save, and run again)')
        print('counts do NOT tie to the Summary sheet: do not upload')
        return 1
    print('counts tie to the Summary sheet (%d checks)' % len(ties))
    return 1 if problems else 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
